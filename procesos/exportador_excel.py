from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from openpyxl.styles import PatternFill


RUTA_PRUEBAS = Path(
    r"C:\Proyectos\Pruebas_Plataforma\Efectivo"
)


def exportar_excel(registros, criterio):

    RUTA_PRUEBAS.mkdir(
        parents=True,
        exist_ok=True
    )

    archivo_salida = (
        RUTA_PRUEBAS /
        f"Cruce_Efectivo_{criterio}.xlsx"
    )

    # -------------------------
    # DataFrames
    # -------------------------

    df = pd.DataFrame(registros)

    if not df.empty:

        df["Cuenta"] = df["Cuenta"].astype(str)

        df["Saldo Extracto"] = (
            pd.to_numeric(
                df["Saldo Extracto"],
                errors="coerce"
            )
            .fillna(0)
            .round(2)
        )

    df_control = pd.DataFrame(
        {
            "Campo": [
                "Filtro",
                "PDF encontrados",
                "PDF procesados"
            ],
            "Valor": [
                criterio,
                len(registros),
                len(registros)
            ]
        }
    )

    # -------------------------
    # Crear Excel
    # -------------------------

    with pd.ExcelWriter(
        archivo_salida,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            sheet_name="Extractos",
            index=False
        )

        df_control.to_excel(
            writer,
            sheet_name="Control",
            index=False
        )

    # -------------------------
    # Formato
    # -------------------------

    wb = load_workbook(archivo_salida)

    color_encabezado = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    fuente_encabezado = Font(
        color="FFFFFF",
        bold=True
    )

    # ==================================
    # HOJA EXTRACTOS
    # ==================================

    ws = wb["Extractos"]

    for fila in ws.iter_rows(
        min_row=1,
        max_row=1
    ):
        for celda in fila:

            celda.fill = color_encabezado
            celda.font = fuente_encabezado

            celda.alignment = Alignment(
                horizontal="center"
            )

    # Formato de datos
    for fila in range(2, ws.max_row + 1):

        # Cuenta
        ws[f"A{fila}"].alignment = Alignment(
            horizontal="left"
        )

        # Saldo
        ws[f"B{fila}"].number_format = "#,##0.00"

        ws[f"B{fila}"].alignment = Alignment(
            horizontal="right"
        )

        # Archivo
        ws[f"C{fila}"].alignment = Alignment(
            horizontal="right"
        )

    # Autofiltro
    ws.auto_filter.ref = ws.dimensions

    # Anchos fijos
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 40

    # ==================================
    # HOJA CONTROL
    # ==================================

    ws_control = wb["Control"]

    for fila in ws_control.iter_rows(
        min_row=1,
        max_row=1
    ):
        for celda in fila:

            celda.fill = color_encabezado
            celda.font = fuente_encabezado

            celda.alignment = Alignment(
                horizontal="center"
            )

    ws_control.column_dimensions["A"].width = 25
    ws_control.column_dimensions["B"].width = 20

    wb.save(archivo_salida)

    return archivo_salida